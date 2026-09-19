from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
import os

def set_font(paragraph, font_name='Times New Roman', font_size=12):
    for run in paragraph.runs:
        run.font.name = font_name
        run.font.size = Pt(font_size)

def add_heading_styled(doc, text, level=1):
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
    return heading

def add_paragraph_styled(doc, text, bold=False):
    p = doc.add_paragraph(text)
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        if bold:
            run.font.bold = True
    return p

def add_code_block(doc, code):
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = 'Courier New'
    run.font.size = Pt(10)
    p.paragraph_format.left_indent = Inches(0.5)
    return p

doc = Document()

style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)

doc.add_paragraph()
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Republic Central Colleges\nAutomated Payroll System')
run.font.name = 'Times New Roman'
run.font.size = Pt(16)
run.font.bold = True

doc.add_paragraph()

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = subtitle.add_run('Data Structure')
run.font.name = 'Times New Roman'
run.font.size = Pt(14)
run.font.bold = True

doc.add_paragraph()
doc.add_paragraph()

sections = [
    {
        'title': '1. Database Overview',
        'content': '''The RCC-APS uses MySQL (MariaDB 10.4) as its database management system. The database is named "rcc_aps" and uses the utf8mb4 character set to support peso signs and emoji characters. The system consists of 16 tables organized to store employee information, payroll records, government contribution tables, leave management, attendance tracking, and system configuration.'''
    },
    {
        'title': '2. Table Descriptions',
        'content': None
    },
    {
        'title': '2.1 users',
        'content': 'Stores employee information including personal details, job classification, salary, and work-related data.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Unique employee identifier'),
            ('first', 'VARCHAR(100)', 'NOT NULL', 'First name'),
            ('last', 'VARCHAR(100)', 'NOT NULL', 'Last name'),
            ('email', 'VARCHAR(200)', 'UNIQUE', 'Email address'),
            ('role', 'VARCHAR(50)', 'NOT NULL', 'User role (Employee, Finance Officer, Admin)'),
            ('status', 'VARCHAR(20)', 'NOT NULL', 'Employment status (Active, Inactive)'),
            ('date', 'VARCHAR(40)', 'NOT NULL', 'Date of birth'),
            ('salary', 'DECIMAL(12,2)', 'NOT NULL', 'Monthly salary'),
            ('classification', 'VARCHAR(50)', 'NOT NULL', 'Teaching or Non-Teaching'),
            ('department', 'VARCHAR(120)', 'NOT NULL', 'Department assignment'),
            ('hireDate', 'VARCHAR(30)', 'NOT NULL', 'Date of employment'),
            ('hourlyRate', 'DECIMAL(10,2)', 'NOT NULL', 'Hourly rate calculation'),
            ('absenceDays', 'DECIMAL(6,2)', 'NOT NULL', 'Number of absent days'),
            ('hoursWorked', 'DECIMAL(8,2)', 'NOT NULL', 'Total hours worked'),
            ('endDate', 'VARCHAR(30)', 'NOT NULL', 'Employment end date')
        ]
    },
    {
        'title': '2.2 auth',
        'content': 'Stores login credentials and role-based access control information for each user.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'User identifier'),
            ('email', 'VARCHAR(200)', 'UNIQUE', 'Login email'),
            ('pass', 'VARCHAR(200)', 'NOT NULL', 'Password (bcrypt hashed)'),
            ('label', 'VARCHAR(50)', 'NOT NULL', 'Display name for role'),
            ('icon', 'VARCHAR(20)', 'NOT NULL', 'Role icon'),
            ('page', 'VARCHAR(50)', 'NOT NULL', 'Assigned dashboard page'),
            ('role', 'VARCHAR(20)', 'NOT NULL', 'Role key (admin, finance, employee)'),
            ('status', 'VARCHAR(20)', 'NOT NULL', 'Account status')
        ]
    },
    {
        'title': '2.3 profiles',
        'content': 'Stores additional contact information for employees.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'User identifier'),
            ('empPhone', 'VARCHAR(40)', 'NOT NULL', 'Employee phone number'),
            ('empAddr', 'VARCHAR(255)', 'NOT NULL', 'Employee address')
        ]
    },
    {
        'title': '2.4 cash_advances',
        'content': 'Tracks cash advance balances for employees.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'User identifier'),
            ('amount', 'DECIMAL(12,2)', 'NOT NULL', 'Cash advance amount')
        ]
    },
    {
        'title': '2.5 bonus_types',
        'content': 'Defines different types of bonuses that can be awarded to employees.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Bonus type identifier'),
            ('name', 'VARCHAR(120)', 'NOT NULL', 'Bonus name'),
            ('kind', 'VARCHAR(30)', 'NOT NULL', 'Bonus category'),
            ('value', 'DECIMAL(10,2)', 'NOT NULL', 'Bonus amount or percentage'),
            ('desc', 'VARCHAR(255)', 'NOT NULL', 'Bonus description'),
            ('status', 'VARCHAR(20)', 'NOT NULL', 'Active or Inactive'),
            ('releaseDate', 'VARCHAR(30)', 'NOT NULL', 'Scheduled release date'),
            ('endDate', 'VARCHAR(30)', 'NOT NULL', 'End validity date')
        ]
    },
    {
        'title': '2.6 bonus_payouts',
        'content': 'Records bonus payouts to employees per year.',
        'columns': [
            ('empId', 'INT', 'Composite PK', 'Employee identifier'),
            ('year', 'INT', 'Composite PK', 'Payout year'),
            ('date', 'VARCHAR(40)', 'NOT NULL', 'Payout date')
        ]
    },
    {
        'title': '2.7 deduction_categories',
        'content': 'Stores additional deduction categories beyond government-mandated contributions.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Deduction category identifier'),
            ('name', 'VARCHAR(120)', 'NOT NULL', 'Deduction name'),
            ('type', 'VARCHAR(30)', 'NOT NULL', 'Deduction type'),
            ('amount', 'DECIMAL(12,2)', 'NOT NULL', 'Deduction amount'),
            ('desc', 'VARCHAR(255)', 'NOT NULL', 'Deduction description'),
            ('status', 'VARCHAR(20)', 'NOT NULL', 'Active or Inactive')
        ]
    },
    {
        'title': '2.8 payroll_records',
        'content': 'Stores monthly payroll computation results for each employee.',
        'columns': [
            ('id', 'INT AUTO_INCREMENT', 'Primary Key', 'Payroll record identifier'),
            ('empId', 'INT', 'NOT NULL', 'Employee identifier'),
            ('name', 'VARCHAR(150)', 'NOT NULL', 'Employee name'),
            ('month', 'INT', 'NOT NULL', 'Payroll month (1-12)'),
            ('monthName', 'VARCHAR(20)', 'NOT NULL', 'Month name'),
            ('year', 'INT', 'NOT NULL', 'Payroll year'),
            ('period', 'VARCHAR(50)', 'NOT NULL', 'Payroll period'),
            ('gross', 'DECIMAL(12,2)', 'NOT NULL', 'Gross pay'),
            ('deMinimisSum', 'DECIMAL(12,2)', 'NOT NULL', 'De minimis benefits'),
            ('thirteenth', 'DECIMAL(12,2)', 'NOT NULL', '13th month pay'),
            ('thirteenTaxable', 'DECIMAL(12,2)', 'NOT NULL', 'Taxable 13th month'),
            ('thirteenExempt', 'DECIMAL(12,2)', 'NOT NULL', 'Exempt 13th month'),
            ('thirteenTax', 'DECIMAL(12,2)', 'NOT NULL', 'Tax on 13th month'),
            ('net', 'DECIMAL(12,2)', 'NOT NULL', 'Net pay'),
            ('sss', 'DECIMAL(12,2)', 'NOT NULL', 'SSS contribution'),
            ('phil', 'DECIMAL(12,2)', 'NOT NULL', 'PhilHealth contribution'),
            ('pag', 'DECIMAL(12,2)', 'NOT NULL', 'Pag-IBIG contribution'),
            ('bir', 'DECIMAL(12,2)', 'NOT NULL', 'BIR withholding tax'),
            ('cash', 'DECIMAL(12,2)', 'NOT NULL', 'Cash advance deduction'),
            ('bonuses', 'VARCHAR(255)', 'NOT NULL', 'Applied bonuses (JSON)'),
            ('date', 'VARCHAR(40)', 'NOT NULL', 'Payroll processing date')
        ]
    },
    {
        'title': '2.9 sss_table',
        'content': 'Stores the Social Security System (SSS) contribution schedule based on salary ranges.',
        'columns': [
            ('id', 'INT AUTO_INCREMENT', 'Primary Key', 'Row identifier'),
            ('range', 'VARCHAR(80)', 'NOT NULL', 'Salary range label'),
            ('emp', 'DECIMAL(10,2)', 'NOT NULL', 'Employee contribution'),
            ('er', 'DECIMAL(10,2)', 'NOT NULL', 'Employer contribution'),
            ('min', 'DECIMAL(12,2)', 'NULL', 'Minimum salary for range'),
            ('max', 'DECIMAL(12,2)', 'NULL', 'Maximum salary for range')
        ]
    },
    {
        'title': '2.10 phil_data',
        'content': 'Stores PhilHealth insurance configuration (single row).',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Always 1'),
            ('rate', 'DECIMAL(6,2)', 'NOT NULL', 'Contribution rate (%)'),
            ('min', 'DECIMAL(12,2)', 'NOT NULL', 'Minimum premium'),
            ('max', 'DECIMAL(12,2)', 'NOT NULL', 'Maximum premium'),
            ('floor', 'DECIMAL(12,2)', 'NOT NULL', 'Salary floor'),
            ('ceiling', 'DECIMAL(12,2)', 'NOT NULL', 'Salary ceiling')
        ]
    },
    {
        'title': '2.11 pag_data',
        'content': 'Stores Pag-IBIG (HDMF) contribution configuration (single row).',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Always 1'),
            ('rate1', 'DECIMAL(6,2)', 'NOT NULL', 'Employee rate (%)'),
            ('rate2', 'DECIMAL(6,2)', 'NOT NULL', 'Employer rate (%)'),
            ('empMax', 'DECIMAL(12,2)', 'NOT NULL', 'Employee max contribution'),
            ('erMax', 'DECIMAL(12,2)', 'NOT NULL', 'Employer max contribution')
        ]
    },
    {
        'title': '2.12 bir_table',
        'content': 'Stores the Bureau of Internal Revenue (BIR) withholding tax table under the TRAIN Law.',
        'columns': [
            ('id', 'INT AUTO_INCREMENT', 'Primary Key', 'Row identifier'),
            ('range', 'VARCHAR(80)', 'NOT NULL', 'Tax bracket label'),
            ('base', 'DECIMAL(12,2)', 'NOT NULL', 'Base tax amount'),
            ('rate', 'DECIMAL(6,2)', 'NOT NULL', 'Tax rate (%)'),
            ('min', 'DECIMAL(12,2)', 'NULL', 'Minimum taxable income'),
            ('max', 'DECIMAL(12,2)', 'NULL', 'Maximum taxable income')
        ]
    },
    {
        'title': '2.13 app_config',
        'content': 'Key-value store for system-wide configuration settings.',
        'columns': [
            ('cfg_key', 'VARCHAR(60)', 'Primary Key', 'Configuration key'),
            ('cfg_value', 'VARCHAR(1024)', 'NOT NULL', 'Configuration value')
        ]
    },
    {
        'title': '2.14 leave_types',
        'content': 'Defines types of leaves available to employees.',
        'columns': [
            ('id', 'INT', 'Primary Key', 'Leave type identifier'),
            ('name', 'VARCHAR(80)', 'NOT NULL', 'Leave type name'),
            ('isPaid', 'TINYINT(1)', 'NOT NULL', 'Paid (1) or Unpaid (0)')
        ]
    },
    {
        'title': '2.15 leave_records',
        'content': 'Stores leave records filed by employees.',
        'columns': [
            ('id', 'INT AUTO_INCREMENT', 'Primary Key', 'Leave record identifier'),
            ('empId', 'INT', 'NOT NULL', 'Employee identifier'),
            ('typeId', 'INT', 'NOT NULL', 'Leave type identifier'),
            ('startDate', 'VARCHAR(20)', 'NOT NULL', 'Leave start date'),
            ('endDate', 'VARCHAR(20)', 'NOT NULL', 'Leave end date'),
            ('days', 'INT', 'NOT NULL', 'Number of leave days')
        ]
    },
    {
        'title': '2.16 attendance',
        'content': 'Stores daily attendance records imported from HRIS.',
        'columns': [
            ('id', 'INT AUTO_INCREMENT', 'Primary Key', 'Attendance record identifier'),
            ('empId', 'INT', 'NOT NULL', 'Internal employee ID'),
            ('employeeId', 'VARCHAR(40)', 'NOT NULL', 'HRIS employee ID'),
            ('employeeEmail', 'VARCHAR(200)', 'NOT NULL', 'Employee email'),
            ('date', 'VARCHAR(20)', 'NOT NULL', 'Attendance date'),
            ('clockInAt', 'VARCHAR(40)', 'NOT NULL', 'Clock-in time'),
            ('clockOutAt', 'VARCHAR(40)', 'NOT NULL', 'Clock-out time'),
            ('hours', 'DECIMAL(6,2)', 'NOT NULL', 'Hours worked')
        ]
    }
]

for section in sections:
    add_heading_styled(doc, section['title'], level=1)
    if section['content']:
        add_paragraph_styled(doc, section['content'])
    
    if 'columns' in section:
        doc.add_paragraph()
        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Column Name'
        hdr_cells[1].text = 'Data Type'
        hdr_cells[2].text = 'Constraint'
        hdr_cells[3].text = 'Description'
        
        for cell in hdr_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = 'Times New Roman'
                    run.font.size = Pt(11)
                    run.font.bold = True
        
        for col_name, data_type, constraint, description in section['columns']:
            row_cells = table.add_row().cells
            row_cells[0].text = col_name
            row_cells[1].text = data_type
            row_cells[2].text = constraint
            row_cells[3].text = description
            
            for cell in row_cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(11)
        
        doc.add_paragraph()

add_heading_styled(doc, '3. SQL CREATE TABLE Statements', level=1)
add_paragraph_styled(doc, 'The following SQL statements define the database schema for the RCC-APS system:')

sql_code = """-- RCC-APS MySQL Schema (MariaDB 10.4)
-- Database: rcc_aps

CREATE DATABASE IF NOT EXISTS rcc_aps
  CHARACTER SET utf8mb4
  COLLATE utf8mb4_unicode_ci;

USE rcc_aps;

-- USERS
CREATE TABLE IF NOT EXISTS users (
  id             INT PRIMARY KEY,
  first          VARCHAR(100) NOT NULL DEFAULT '',
  last           VARCHAR(100) NOT NULL DEFAULT '',
  email          VARCHAR(200) NOT NULL DEFAULT '',
  role           VARCHAR(50)  NOT NULL DEFAULT 'Employee',
  status         VARCHAR(20)  NOT NULL DEFAULT 'Active',
  date           VARCHAR(40)  NOT NULL DEFAULT '',
  salary         DECIMAL(12,2) NOT NULL DEFAULT 0,
  classification VARCHAR(50)  NOT NULL DEFAULT 'Non-Teaching',
  department     VARCHAR(120) NOT NULL DEFAULT '',
  hireDate       VARCHAR(30)  NOT NULL DEFAULT '',
  hourlyRate     DECIMAL(10,2) NOT NULL DEFAULT 0,
  absenceDays    DECIMAL(6,2) NOT NULL DEFAULT 0,
  hoursWorked    DECIMAL(8,2) NOT NULL DEFAULT 0,
  endDate        VARCHAR(30)  NOT NULL DEFAULT '',
  UNIQUE KEY uk_users_email (email)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- AUTH
CREATE TABLE IF NOT EXISTS auth (
  id      INT PRIMARY KEY,
  email   VARCHAR(200) NOT NULL DEFAULT '',
  pass    VARCHAR(200) NOT NULL DEFAULT '',
  label   VARCHAR(50)  NOT NULL DEFAULT '',
  icon    VARCHAR(20)  NOT NULL DEFAULT '',
  page    VARCHAR(50)  NOT NULL DEFAULT '',
  role    VARCHAR(20)  NOT NULL DEFAULT 'employee',
  status  VARCHAR(20)  NOT NULL DEFAULT 'Active',
  UNIQUE KEY uk_auth_email (email)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- PROFILES
CREATE TABLE IF NOT EXISTS profiles (
  id       INT PRIMARY KEY,
  empPhone VARCHAR(40) NOT NULL DEFAULT '',
  empAddr  VARCHAR(255) NOT NULL DEFAULT ''
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- CASH ADVANCES
CREATE TABLE IF NOT EXISTS cash_advances (
  id     INT PRIMARY KEY,
  amount DECIMAL(12,2) NOT NULL DEFAULT 0
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- BONUS TYPES
CREATE TABLE IF NOT EXISTS bonus_types (
  id          INT PRIMARY KEY,
  name        VARCHAR(120) NOT NULL DEFAULT '',
  kind        VARCHAR(30)  NOT NULL DEFAULT '',
  value       DECIMAL(10,2) NOT NULL DEFAULT 0,
  `desc`      VARCHAR(255) NOT NULL DEFAULT '',
  status      VARCHAR(20)  NOT NULL DEFAULT 'Active',
  releaseDate VARCHAR(30)  NOT NULL DEFAULT '',
  endDate     VARCHAR(30)  NOT NULL DEFAULT ''
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- BONUS PAYOUTS
CREATE TABLE IF NOT EXISTS bonus_payouts (
  empId INT NOT NULL,
  year  INT NOT NULL DEFAULT 0,
  date  VARCHAR(40) NOT NULL DEFAULT '',
  PRIMARY KEY (empId, year)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- DEDUCTION CATEGORIES
CREATE TABLE IF NOT EXISTS deduction_categories (
  id     INT PRIMARY KEY,
  name   VARCHAR(120) NOT NULL DEFAULT '',
  type   VARCHAR(30)  NOT NULL DEFAULT '',
  amount DECIMAL(12,2) NOT NULL DEFAULT 0,
  `desc` VARCHAR(255) NOT NULL DEFAULT '',
  status VARCHAR(20)  NOT NULL DEFAULT 'Active'
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- PAYROLL RECORDS
CREATE TABLE IF NOT EXISTS payroll_records (
  id            INT AUTO_INCREMENT PRIMARY KEY,
  empId         INT NOT NULL,
  name          VARCHAR(150) NOT NULL DEFAULT '',
  month         INT NOT NULL DEFAULT 0,
  monthName     VARCHAR(20)  NOT NULL DEFAULT '',
  year          INT NOT NULL DEFAULT 0,
  period        VARCHAR(50)  NOT NULL DEFAULT '',
  gross         DECIMAL(12,2) NOT NULL DEFAULT 0,
  deMinimisSum  DECIMAL(12,2) NOT NULL DEFAULT 0,
  thirteenth    DECIMAL(12,2) NOT NULL DEFAULT 0,
  thirteenTaxable DECIMAL(12,2) NOT NULL DEFAULT 0,
  thirteenExempt  DECIMAL(12,2) NOT NULL DEFAULT 0,
  thirteenTax     DECIMAL(12,2) NOT NULL DEFAULT 0,
  net           DECIMAL(12,2) NOT NULL DEFAULT 0,
  sss           DECIMAL(12,2) NOT NULL DEFAULT 0,
  phil          DECIMAL(12,2) NOT NULL DEFAULT 0,
  pag           DECIMAL(12,2) NOT NULL DEFAULT 0,
  bir           DECIMAL(12,2) NOT NULL DEFAULT 0,
  cash          DECIMAL(12,2) NOT NULL DEFAULT 0,
  bonuses       VARCHAR(255) NOT NULL DEFAULT '[]',
  date          VARCHAR(40)  NOT NULL DEFAULT '',
  KEY idx_payroll_emp (empId),
  KEY idx_payroll_period (year, month)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- SSS CONTRIBUTION TABLE
CREATE TABLE IF NOT EXISTS sss_table (
  id    INT AUTO_INCREMENT PRIMARY KEY,
  `range` VARCHAR(80) NOT NULL DEFAULT '',
  emp   DECIMAL(10,2) NOT NULL DEFAULT 0,
  er    DECIMAL(10,2) NOT NULL DEFAULT 0,
  min   DECIMAL(12,2) NULL,
  max   DECIMAL(12,2) NULL
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- PHILHEALTH CONFIG
CREATE TABLE IF NOT EXISTS phil_data (
  id      INT PRIMARY KEY DEFAULT 1,
  rate    DECIMAL(6,2) NOT NULL DEFAULT 5,
  min     DECIMAL(12,2) NOT NULL DEFAULT 500,
  max     DECIMAL(12,2) NOT NULL DEFAULT 5000,
  floor   DECIMAL(12,2) NOT NULL DEFAULT 10000,
  ceiling DECIMAL(12,2) NOT NULL DEFAULT 100000
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- PAG-IBIG CONFIG
CREATE TABLE IF NOT EXISTS pag_data (
  id     INT PRIMARY KEY DEFAULT 1,
  rate1  DECIMAL(6,2) NOT NULL DEFAULT 1,
  rate2  DECIMAL(6,2) NOT NULL DEFAULT 2,
  empMax DECIMAL(12,2) NOT NULL DEFAULT 100,
  erMax  DECIMAL(12,2) NOT NULL DEFAULT 100
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- BIR WITHHOLDING TABLE
CREATE TABLE IF NOT EXISTS bir_table (
  id    INT AUTO_INCREMENT PRIMARY KEY,
  `range` VARCHAR(80)  NOT NULL DEFAULT '',
  base  DECIMAL(12,2) NOT NULL DEFAULT 0,
  rate  DECIMAL(6,2)  NOT NULL DEFAULT 0,
  min   DECIMAL(12,2) NULL,
  max   DECIMAL(12,2) NULL
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- APP CONFIG
CREATE TABLE IF NOT EXISTS app_config (
  cfg_key   VARCHAR(60) PRIMARY KEY,
  cfg_value VARCHAR(1024) NOT NULL
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- LEAVE TYPES
CREATE TABLE IF NOT EXISTS leave_types (
  id     INT PRIMARY KEY,
  name   VARCHAR(80) NOT NULL DEFAULT '',
  isPaid TINYINT(1)  NOT NULL DEFAULT 1
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- LEAVE RECORDS
CREATE TABLE IF NOT EXISTS leave_records (
  id        INT AUTO_INCREMENT PRIMARY KEY,
  empId     INT NOT NULL,
  typeId    INT NOT NULL,
  startDate VARCHAR(20) NOT NULL DEFAULT '',
  endDate   VARCHAR(20) NOT NULL DEFAULT '',
  days      INT NOT NULL DEFAULT 0
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;

-- ATTENDANCE
CREATE TABLE IF NOT EXISTS attendance (
  id            INT AUTO_INCREMENT PRIMARY KEY,
  empId         INT NOT NULL DEFAULT 0,
  employeeId    VARCHAR(40) NOT NULL DEFAULT '',
  employeeEmail VARCHAR(200) NOT NULL DEFAULT '',
  date          VARCHAR(20) NOT NULL DEFAULT '',
  clockInAt     VARCHAR(40) NOT NULL DEFAULT '',
  clockOutAt    VARCHAR(40) NOT NULL DEFAULT '',
  hours         DECIMAL(6,2) NOT NULL DEFAULT 0,
  UNIQUE KEY uk_att_emp_date (empId, date)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;"""

add_code_block(doc, sql_code)

doc.add_paragraph()
add_heading_styled(doc, '4. Entity Relationships', level=1)
add_paragraph_styled(doc, 'The following relationships exist between tables in the RCC-APS database:')
doc.add_paragraph()

relationships = [
    'users.id ← auth.id : One-to-one (login credentials)',
    'users.id ← profiles.id : One-to-one (contact information)',
    'users.id ← cash_advances.id : One-to-one (cash advance balance)',
    'users.id → payroll_records.empId : One-to-many (monthly payroll records)',
    'users.id → attendance.empId : One-to-many (daily attendance)',
    'users.id → leave_records.empId : One-to-many (leave records)',
    'leave_records.typeId → leave_types.id : Many-to-one (leave type reference)',
    'bonus_payouts.empId → users.id : Many-to-one (bonus release)',
    'payroll_records.empId → users.id : Many-to-one (payroll computation)'
]

for rel in relationships:
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(rel)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)

output_path = r'C:\Users\Admin\Pictures\RCC-APS SYSTEM\RCC-APS_Data_Structure.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
